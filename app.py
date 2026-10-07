import os
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import tempfile

# ============================================================
# CONFIGURATION DE LA PAGE WEB
# ============================================================
st.set_page_config(page_title="Générateur Organigramme PPTX", layout="wide")
st.title("🏗️ Générateur d'Organigramme de Chantier")

# ============================================================
# VOS FONCTIONS PYTHON ORIGINALES (Inchangées)
# ============================================================
COLOR_WHITE     = RGBColor(255, 255, 255)
COLOR_RED       = RGBColor(205, 33, 39)
COLOR_SAND      = RGBColor(240, 236, 228)
COLOR_GREEN     = RGBColor(46, 168, 139)
COLOR_TEXT_DARK = RGBColor(112, 112, 112)
COLOR_BLUE      = RGBColor(85, 198, 221)

def ajouter_zone_image_ronde_cliquable(slide, x, y, diametre=Inches(1.50), chemin_placeholder="placeholder_rond.png"):
    if not os.path.exists(chemin_placeholder):
        from PIL import Image, ImageDraw
        img = Image.new("RGB", (400, 400), "#FFFFFF")
        draw = ImageDraw.Draw(img)
        draw.ellipse([(10, 10), (390, 390)], fill="#F8F8F8", outline="#D2D2D2", width=3)
        img.save(chemin_placeholder)
    zone_ronde = slide.shapes.add_picture(chemin_placeholder, x, y, width=diametre, height=diametre)
    zone_ronde.auto_shape_type = MSO_SHAPE.OVAL
    return zone_ronde

def ajouter_titre_section_pptx(slide, texte, x, y, largeur, hauteur, couleur_fond, couleur_texte):
    diametre = Inches(0.45)
    x_rectangle = x + (diametre / 2)
    largeur_rectangle = largeur - (diametre / 2)

    rectangle = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_rectangle, y, largeur_rectangle, hauteur)
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = couleur_fond
    rectangle.line.fill.background()

    tf = rectangle.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)

    p = tf.paragraphs[0]
    p.text = texte
    p.font.color.rgb = couleur_texte
    p.font.size = Pt(11)
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

    y_rond = y + (hauteur - diametre) / 2
    macaron = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y_rond, diametre, diametre)
    macaron.fill.solid()
    macaron.shadow.inherit = False

    texte_clean = texte.upper()
    nom_logo_cible = None

    if "MOE" in texte_clean or "OEUVRE" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_RED
        nom_logo_cible = "Machine"
    elif "PILOTAGE" in texte_clean or "COPIL" in texte_clean or "CHANTIER" in texte_clean or "DIRECTION" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_RED
        nom_logo_cible = "Engrenage"
    elif "PRODUCTION" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_RED
        nom_logo_cible = "Grue"
    elif "TRAITANT" in texte_clean or "CO-" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_BLUE
        nom_logo_cible = "Grue"
    elif "MAINTENEUR" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_TEXT_DARK
        nom_logo_cible = "Grue"
    elif "SUPPORT" in texte_clean or "POLE" in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_RED
        nom_logo_cible = "Main"
    elif "COND." in texte_clean:
        macaron.fill.fore_color.rgb = COLOR_TEXT_DARK
        nom_logo_cible = "Grue"
    else:
        macaron.fill.fore_color.rgb = COLOR_RED

    macaron.line.fill.background()
    
    if nom_logo_cible:
        extensions = [".png", ".PNG", ".jpg", ".JPG", ".jpeg", ".JPEG"]
        chemin_logo = None
        for ext in extensions:
            if os.path.exists(f"{nom_logo_cible}{ext}"):
                chemin_logo = f"{nom_logo_cible}{ext}"
                break

        if chemin_logo:
            if nom_logo_cible == "Machine":
                taille_logo = Inches(0.45)
                logo_shape = slide.shapes.add_picture(chemin_logo, x, y_rond, width=taille_logo, height=taille_logo)
                logo_shape.auto_shape_type = MSO_SHAPE.OVAL
            else:
                taille_logo = Inches(0.30)
                x_logo = x + (diametre - taille_logo) / 2
                y_logo = y_rond + (diametre - taille_logo) / 2
                slide.shapes.add_picture(chemin_logo, x_logo, y_logo, width=taille_logo, height=taille_logo)

def ajouter_fiche_personne_pptx(slide, role, name, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_SAND
    card.line.color.rgb = COLOR_SAND
    card.line.width = Pt(0.5)

    red_bar = slide.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, left - Inches(0.01), top, width + Inches(0.02), Inches(0.13))
    red_bar.adjustments[0] = 0.25
    red_bar.fill.solid()
    red_bar.fill.fore_color.rgb = COLOR_RED
    red_bar.line.fill.background()
    red_bar.shadow.inherit = False

    text_box = slide.shapes.add_textbox(left + Inches(0.02), top + Inches(0.14), width - Inches(0.04), height - Inches(0.16))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p_role = tf.paragraphs[0]
    p_role.text = role
    p_role.font.name = "Poppins"
    p_role.font.size = Pt(8)
    p_role.font.bold = True
    p_role.font.color.rgb = COLOR_TEXT_DARK
    p_role.alignment = PP_ALIGN.CENTER

    p_name = tf.add_paragraph()
    p_name.text = name
    p_name.font.name = "Poppins"
    p_name.font.size = Pt(7.5)
    p_name.font.color.rgb = COLOR_TEXT_DARK
    p_name.alignment = PP_ALIGN.CENTER

    diametre_rond = Inches(0.59)
    x_rond = left + (width / 2) - (diametre_rond / 2)
    y_rond = top - (diametre_rond) + Inches(0.13)
    ajouter_zone_image_ronde_cliquable(slide, x_rond, y_rond, diametre_rond)

def ajouter_fiche_personne_verte_pptx(slide, role, name, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_SAND
    card.line.color.rgb = COLOR_SAND

    gre_bar = slide.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, left - Inches(0.01), top, width + Inches(0.02), Inches(0.13))
    gre_bar.adjustments[0] = 0.25
    gre_bar.fill.solid()
    gre_bar.fill.fore_color.rgb = COLOR_GREEN
    gre_bar.line.fill.background()

    text_box = slide.shapes.add_textbox(left + Inches(0.02), top + Inches(0.14), width - Inches(0.04), height - Inches(0.16))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    intitule_role = role
    nom_entreprise = ""
    if " (" in role:
        parts = role.split(" (")
        intitule_role = parts[0]
        nom_entreprise = parts[1].replace(")", "")

    p_role = tf.paragraphs[0]
    p_role.text = intitule_role
    p_role.font.size = Pt(7.5)
    p_role.font.bold = True
    p_role.alignment = PP_ALIGN.CENTER

    p_ent = tf.add_paragraph()
    p_ent.text = nom_entreprise
    p_ent.font.size = Pt(7)
    p_ent.alignment = PP_ALIGN.CENTER

    p_name = tf.add_paragraph()
    p_name.text = name
    p_name.font.size = Pt(7)
    p_name.alignment = PP_ALIGN.CENTER

    largeur_img = Inches(1.20)
    hauteur_img = Inches(0.50)
    x_img = left + (width / 2) - (largeur_img / 2)
    y_img = top - hauteur_img
    ajouter_zone_image_cliquable(slide, x_img, y_img, largeur_img, hauteur_img)

def ajouter_fiche_personne_gris_pptx(slide, role, name, left, top, width, height):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_SAND
    card.line.color.rgb = COLOR_SAND

    gri_bar = slide.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, left - Inches(0.01), top, width + Inches(0.02), Inches(0.13))
    gri_bar.adjustments[0] = 0.25
    gri_bar.fill.solid()
    gri_bar.fill.fore_color.rgb = COLOR_TEXT_DARK
    gri_bar.line.fill.background()

    text_box = slide.shapes.add_textbox(left + Inches(0.02), top + Inches(0.14), width - Inches(0.04), height - Inches(0.16))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p_role = tf.paragraphs[0]
    p_role.text = role
    p_role.font.size = Pt(8)
    p_role.font.bold = True
    p_role.alignment = PP_ALIGN.CENTER

    p_name = tf.add_paragraph()
    p_name.text = name
    p_name.font.size = Pt(7.5)
    p_name.alignment = PP_ALIGN.CENTER

    largeur_img = Inches(1.20)
    hauteur_img = Inches(0.50)
    x_img = left + (width / 2) - (largeur_img / 2)
    y_img = top - hauteur_img
    ajouter_zone_image_cliquable(slide, x_img, y_img, largeur_img, hauteur_img)

def ajouter_zone_image_cliquable(slide, x, y, largeur, hauteur, chemin_placeholder="placeholder_temp.png"):
    if not isinstance(chemin_placeholder, str) or not os.path.exists(chemin_placeholder):
        nom_sauvegarde = "placeholder_temp.png"
        if not os.path.exists(nom_sauvegarde):
            from PIL import Image, ImageDraw
            img = Image.new("RGB", (300, 200), "#F8F8F8")
            draw = ImageDraw.Draw(img)
            draw.rectangle([(0, 0), (299, 199)], outline="#D2D2D2", width=2)
            img.save(nom_sauvegarde)
        chemin_placeholder = nom_sauvegarde

    return slide.shapes.add_picture(chemin_placeholder, x, y, width=largeur, height=hauteur)

def generer_presentation(donnees):
    prs = Presentation()
    LARGEUR_PAGE = Inches(14.698)
    HAUTEUR_PAGE = Inches(8.267)
    CENTRE_PAGE = LARGEUR_PAGE / 2

    prs.slide_width = LARGEUR_PAGE
    prs.slide_height = HAUTEUR_PAGE
    blank_layout = prs.slide_layouts[6] 

    fiche_w = Inches(1.39)
    fiche_h = Inches(0.82)
    y_sections_titres_bas = Inches(3.85)
    y_fiches_moe = Inches(4.93)

    # --- SLIDE 1 : CONCEPTION ---
    slide1 = prs.slides.add_slide(blank_layout)
    
    chemin_logo_lg = "Lg.png" if os.path.exists("Lg.png") else "placeholder_temp.png"
    ajouter_zone_image_cliquable(slide1, Inches(4.35), Inches(0.15), Inches(0.85), Inches(0.70), chemin_placeholder=chemin_logo_lg)

    title_box1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, CENTRE_PAGE - Inches(2), Inches(0.20), Inches(4), Inches(0.60))
    title_box1.fill.solid()
    title_box1.fill.fore_color.rgb = COLOR_SAND
    title_box1.line.color.rgb = RGBColor(210, 210, 210)
    p_t1 = title_box1.text_frame.paragraphs[0]
    p_t1.text = "ORGANIGRAMME PHASE CONCEPTION"
    p_t1.font.size = Pt(14)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_RED
    p_t1.alignment = PP_ALIGN.CENTER

    # MOA
    gauche_moa = Inches(0.40)
    largeur_moa = Inches(2.20)
    ajouter_zone_image_cliquable(slide1, gauche_moa, Inches(0.10), largeur_moa, Inches(1.00))
    
    box_nom_moa1 = slide1.shapes.add_textbox(gauche_moa, Inches(1.10), largeur_moa, Inches(0.35))
    p_nom_moa1 = box_nom_moa1.text_frame.paragraphs[0]
    p_nom_moa1.text = donnees["moa"].upper()
    p_nom_moa1.font.size = Pt(12)
    p_nom_moa1.font.bold = True
    p_nom_moa1.alignment = PP_ALIGN.CENTER

    forme1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, gauche_moa, Inches(1.50), largeur_moa, Inches(0.80))
    forme1.fill.solid()
    forme1.fill.fore_color.rgb = RGBColor(220, 230, 242)
    forme1.line.color.rgb = RGBColor(141, 180, 226)
    tf_m1 = forme1.text_frame
    tf_m1.paragraphs[0].text = "Maîtrise d'Ouvrage"
    tf_m1.paragraphs[0].font.size = Pt(14)
    tf_m1.paragraphs[0].font.bold = True
    p_m2 = tf_m1.add_paragraph()
    p_m2.text = "AMO, CT, CSPS"
    p_m2.font.size = Pt(8)
    p_m2.alignment = PP_ALIGN.CENTER

    # CO-TRAITANT
    if donnees["co_traitant"] and donnees["co_traitant"] not in ["Aucun", ""]:
        ajouter_titre_section_pptx(slide1, "CO-TRAITANT", Inches(3.20), Inches(1.45), Inches(2.40), Inches(0.45), COLOR_SAND, COLOR_BLUE)
        ajouter_zone_image_cliquable(slide1, Inches(3.20) + (Inches(2.40)/2) - (Inches(1.35)/2), Inches(2.05), Inches(1.35), Inches(0.75))

    # DIRECTION
    ajouter_titre_section_pptx(slide1, "DIRECTION DE PROJET", CENTRE_PAGE - Inches(1.2), Inches(1.45), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    ajouter_fiche_personne_pptx(slide1, donnees["direction"]["poste"], donnees["direction"]["nom"], CENTRE_PAGE - (fiche_w/2), Inches(2.55), fiche_w, fiche_h)

    # COPIL
    ajouter_titre_section_pptx(slide1, "COMITÉ DE PILOTAGE", Inches(11.85), Inches(0.15), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    coords_copil = [(Inches(11.65), Inches(1.23)), (Inches(13.15), Inches(1.23)), (Inches(11.65), Inches(2.68)), (Inches(13.15), Inches(2.68))]
    for i, membre in enumerate(donnees["copil"]):
        if i < len(coords_copil):
            ajouter_fiche_personne_pptx(slide1, membre["poste"], membre["nom"], coords_copil[i][0], coords_copil[i][1], fiche_w, fiche_h)

    # MOE
    ajouter_titre_section_pptx(slide1, "Maîtrise d'oeuvre", Inches(1.295), y_sections_titres_bas, Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_GREEN)
    for idx, moe in enumerate(donnees["moe"]):
        x = Inches(0.50) + (idx % 3) * Inches(1.75)
        y = y_fiches_moe + (idx // 3) * Inches(1.45)
        ajouter_fiche_personne_verte_pptx(slide1, moe["role"], moe["nom"], x, y, fiche_w, fiche_h)

    # MAINTENEUR
    if donnees["mainteneur"]["entreprise"] and donnees["mainteneur"]["entreprise"] != "Aucun":
        ajouter_titre_section_pptx(slide1, "Mainteneur", CENTRE_PAGE - Inches(1.2), y_sections_titres_bas, Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_TEXT_DARK)
        ajouter_fiche_personne_gris_pptx(slide1, donnees["mainteneur"]["entreprise"], donnees["mainteneur"]["contact"], CENTRE_PAGE - (fiche_w/2), Inches(4.93), fiche_w, fiche_h)

    # SUPPORT
    ajouter_titre_section_pptx(slide1, "POLE SUPPORT", Inches(9.88), y_sections_titres_bas, Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    coords_services = [(Inches(9.58), Inches(4.93)), (Inches(11.18), Inches(4.93)), (Inches(10.38), Inches(6.38))]
    for idx, service in enumerate(donnees["services_internes"]):
        if idx < len(coords_services):
            ajouter_fiche_personne_pptx(slide1, service["poste"], service["nom"], coords_services[idx][0], coords_services[idx][1], fiche_w, fiche_h)

    # --- SLIDE 2 : REALISATION ---
    slide2 = prs.slides.add_slide(blank_layout)
    ajouter_zone_image_cliquable(slide2, Inches(4.35), Inches(0.15), Inches(0.85), Inches(0.70), chemin_placeholder=chemin_logo_lg)

    title_box2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, CENTRE_PAGE -Inches(2), Inches(0.20), Inches(4), Inches(0.60))
    title_box2.fill.solid()
    title_box2.fill.fore_color.rgb = COLOR_SAND
    title_box2.line.color.rgb = RGBColor(210, 210, 210)
    p_t2 = title_box2.text_frame.paragraphs[0]
    p_t2.text = "ORGANIGRAMME PHASE RÉALISATION"
    p_t2.font.size = Pt(14)
    p_t2.font.bold = True
    p_t2.font.color.rgb = COLOR_RED
    p_t2.alignment = PP_ALIGN.CENTER

    # MOA
    ajouter_zone_image_cliquable(slide2, gauche_moa, Inches(0.10), largeur_moa, Inches(1.00))
    box_nom_moa2 = slide2.shapes.add_textbox(gauche_moa, Inches(1.10), largeur_moa, Inches(0.35))
    box_nom_moa2.text_frame.paragraphs[0].text = donnees["moa"].upper()
    box_nom_moa2.text_frame.paragraphs[0].font.size = Pt(12)
    box_nom_moa2.text_frame.paragraphs[0].font.bold = True
    box_nom_moa2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    forme2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, gauche_moa, Inches(1.50), largeur_moa, Inches(0.80))
    forme2.fill.solid()
    forme2.fill.fore_color.rgb = RGBColor(220, 230, 242)
    forme2.line.color.rgb = RGBColor(141, 180, 226)
    tf_m2 = forme2.text_frame
    tf_m2.paragraphs[0].text = "Maîtrise d'Ouvrage"
    tf_m2.paragraphs[0].font.size = Pt(14)
    tf_m2.paragraphs[0].font.bold = True
    p_m2_2 = tf_m2.add_paragraph()
    p_m2_2.text = "AMO, CT, CSPS"
    p_m2_2.font.size = Pt(8)
    p_m2_2.alignment = PP_ALIGN.CENTER

    # COPIL S2
    ajouter_titre_section_pptx(slide2, "COMITÉ DE PILOTAGE", Inches(11.85), Inches(0.15), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    for i, membre in enumerate(donnees["copil"]):
        if i < len(coords_copil):
            ajouter_fiche_personne_pptx(slide2, membre["poste"], membre["nom"], coords_copil[i][0], coords_copil[i][1], fiche_w, fiche_h)

    # DIRECTION DE PROJET
    ajouter_titre_section_pptx(slide2, "DIRECTION DE PROJET", CENTRE_PAGE - (Inches(2.4)/2), Inches(1.45), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    ajouter_fiche_personne_pptx(slide2, "Directeur de projet", donnees["direction"]["nom"], CENTRE_PAGE - (fiche_w/2), Inches(2.55), fiche_w, fiche_h)

    # DIRECTEUR CHANTIER
    ajouter_titre_section_pptx(slide2, "DIRECTEUR DE CHANTIER", CENTRE_PAGE - (Inches(2.4)/2), Inches(3.85), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    ajouter_fiche_personne_pptx(slide2, "Dir. Chantier", donnees["directeur_chantier"], CENTRE_PAGE - (fiche_w/2), Inches(4.96), fiche_w, fiche_h)

    # MOE SIMPLIFIEE
    ajouter_titre_section_pptx(slide2, "Pôle MOE", Inches(0.40), Inches(3.25), Inches(3.20), Inches(0.45), COLOR_SAND, COLOR_GREEN)
    cadre_moe = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.40), Inches(3.85), Inches(3.20), Inches(1.00))
    cadre_moe.fill.solid()
    cadre_moe.fill.fore_color.rgb = COLOR_WHITE
    cadre_moe.line.color.rgb = RGBColor(180, 180, 180)
    tf_cadre = cadre_moe.text_frame
    tf_cadre.text = "Équipe détaillée en Phase Conception"
    tf_cadre.paragraphs[0].font.size = Pt(10)
    tf_cadre.paragraphs[0].font.color.rgb = COLOR_TEXT_DARK
    tf_cadre.paragraphs[0].alignment = PP_ALIGN.CENTER

    # CONDUCTEURS TRAVAUX
    conducteurs = donnees.get("conducteurs_travaux", [])[:3]
    if len(conducteurs) > 0:
        largeur_bloc = Inches(3.00)
        espace_horizontal = Inches(0.35)
        X_DEBUT = Inches(4.00)
        X_FIN = Inches(14.20)
        CENTRE_ZONE = X_DEBUT + ((X_FIN - X_DEBUT) / 2)
        largeur_totale = (len(conducteurs) * largeur_bloc) + ((len(conducteurs) - 1) * espace_horizontal)
        x_depart = CENTRE_ZONE - (largeur_totale / 2)

        for idx, cond in enumerate(conducteurs):
            x_colonne = x_depart + idx * (largeur_bloc + espace_horizontal)
            ajouter_titre_section_pptx(slide2, f"PRODUCTION - {cond['secteur'].upper()}", x_colonne, Inches(3.85), largeur_bloc, Inches(0.45), COLOR_SAND, COLOR_RED)
            ajouter_fiche_personne_pptx(slide2, "CTX Principal", cond["nom"], x_colonne + (largeur_bloc / 2) - (fiche_w / 2), Inches(4.96), fiche_w, fiche_h)

    # POLE SUPPORT S2
    ajouter_titre_section_pptx(slide2, "POLE SUPPORT", CENTRE_PAGE - (Inches(2.4) / 2), Inches(6.20), Inches(2.4), Inches(0.45), COLOR_SAND, COLOR_RED)
    x_origine_s2 = Inches(5.44)
    coords_support_s2 = [x_origine_s2, x_origine_s2 + Inches(1.60), x_origine_s2 + Inches(3.20)]
    for idx, service in enumerate(donnees["services_internes"]):
        if idx < len(coords_support_s2):
            ajouter_fiche_personne_pptx(slide2, service["poste"], service["nom"], coords_support_s2[idx], Inches(7.28), fiche_w, fiche_h)

    path_final = "organigramme_final.pptx"
    prs.save(path_final)
    return path_final

# ============================================================
# INTERFACE WEB STREAMLIT
# ============================================================
def parser_texte_liste(texte, cls_role="poste", cls_nom="nom"):
    lignes = texte.split('\n')
    resultats = []
    for ligne in lignes:
        if "/" in ligne:
            nom, role = ligne.split('/', 1)
            resultats.append({cls_nom: nom.strip(), cls_role: role.strip()})
    return resultats

col1, col2 = st.columns([1, 2])

with col1:
    st.header("1. Importer les images")
    st.info("Insérez vos logos (Lg.png, Machine.png, Engrenage.png...) et photos.")
    fichiers_images = st.file_uploader("Glissez vos fichiers ici", accept_multiple_files=True, type=['png', 'jpg', 'jpeg'])
    
    if fichiers_images:
        for f in fichiers_images:
            # Sauvegarde temporaire des images uploadées pour que pptx les lise
            with open(f.name, "wb") as f_out:
                f_out.write(f.getbuffer())

    st.header("2. Remplir les informations")
    moa = st.text_input("Maître d'Ouvrage (MOA)", "Finances Publiques")
    dir_projet = st.text_input("Directeur de Projet (Prénom Nom)", "Alexandre MARTIN")
    
    copil = st.text_area("Comité de Pilotage (Nom / Rôle)", "Rémi HOVAERE / Directeur National\nJean-Stéphane DIDIER / DGA\nCharlotte VIGUIER / Dir. Grands Projets\nMicaël GONCALVES / Dir. Excellence")
    
    moe = st.text_area("Pôle MOE (Nom / Entreprise)", "Sophie VALENTIN / Architecte (RDA Architecture)\nPierre DUBOIS / BE (BET Structure)")
    
    co_traitant = st.text_input("Co-traitant", "Eiffage Construction")
    
    support = st.text_area("Pôle Support (Nom / Rôle)", "Emmanuel SAURIN / Réf. Bas Carbone\nJérôme TRANCHANT / QSE Sécurité\nJérôme JUNIQUE / Chef de service Méthode")
    
    maint_contact = st.text_input("Mainteneur (Contact)", "Thomas ROUSSEL")
    maint_ent = st.text_input("Mainteneur (Entreprise)", "Dalkia")
    
    st.subheader("Uniquement Phase Réalisation")
    dir_chantier = st.text_input("Directeur de Chantier", "Marc DURAND")
    conducteurs = st.text_area("Production / Conducteurs (Nom / Zone)", "Julien FAURE / Bâtiment A\nÉlodie MICHEL / Infrastructures")

with col2:
    st.header("3. Générer le fichier")
    st.write("Cliquez sur le bouton pour compiler les informations et générer votre PowerPoint exactement avec les macros d'origine.")
    
    if st.button("Générer mon Organigramme PPTX", type="primary"):
        # Formatage des données
        donnees = {
            "moa": moa,
            "direction": {"nom": dir_projet, "poste": "Directeur de Projet"},
            "copil": parser_texte_liste(copil),
            "moe": parser_texte_liste(moe, cls_role="role"),
            "co_traitant": co_traitant,
            "services_internes": parser_texte_liste(support),
            "mainteneur": {"contact": maint_contact, "entreprise": maint_ent},
            "directeur_chantier": dir_chantier,
            "conducteurs_travaux": parser_texte_liste(conducteurs, cls_role="secteur")
        }

        # Génération
        with st.spinner("Génération du PowerPoint en cours..."):
            fichier_pptx = generer_presentation(donnees)
            
            with open(fichier_pptx, "rb") as f:
                st.download_button(
                    label="📥 Télécharger le fichier .pptx",
                    data=f,
                    file_name="Organigramme_Chantier.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )